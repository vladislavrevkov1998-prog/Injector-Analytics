"""XLSX import, multi-cell arc entry and calculation/geometry separation."""
from copy import deepcopy
from pathlib import Path
import math
import tempfile
import unittest
from unittest.mock import patch
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QDialog,QDialogButtonBox
from openpyxl import Workbook
from fuelfield.chamber import Primitive
from fuelfield.chamber_import import read_point_tables,parse_points
from fuelfield.chamber_import_gui import PointImportDialog
from fuelfield.core import project_to_json
import test_v06_gui as previous


def workbook(path):
    book=Workbook();sheet=book.active;sheet.title='Точки';sheet.append(['z, мм','R, мм']);sheet.append([1,250]);sheet.cell(5,1,2.125);sheet.cell(5,2,'251,375')
    sheet=book.create_sheet('Другой лист');sheet.append(['R','x']);sheet.append([25,.5]);sheet.append([26,1.5]);book.save(path)


class TestPointTables(unittest.TestCase):
    def test_real_xlsx_rows_units_and_blank_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'точки.xlsx';workbook(path);tables=read_point_tables(path)
            self.assertEqual(parse_points(tables[0][1],0,1),[(2,1.,250.),(5,2.125,251.375)])
            self.assertEqual(parse_points(tables[1][1],1,0,unit='cm'),[(2,5.,250.),(3,15.,260.)])

    def test_invalid_cells_have_exact_excel_row_and_no_partial_result(self):
        for value in (None,True,'NaN','Inf','#DIV/0!','text'):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError,'Строка Excel 7'):
                    parse_points([(2,[1,250]),(7,[value,251])],0,1)
        with self.assertRaisesRegex(ValueError,'Повторная'):parse_points([(2,[1,250]),(3,[1,250])],0,1)

    def test_formula_without_cached_value_is_reported_not_skipped(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'formula.xlsx';book=Workbook();s=book.active;s.append(['z','R']);s.cell(8,1,'=1+2');s.cell(8,2,250);book.save(path)
            with self.assertRaisesRegex(ValueError,'Строка Excel 8'):parse_points(read_point_tables(path)[0][1],0,1)

    def test_bad_workbook_is_a_readable_value_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'broken.xlsx';path.write_text('broken')
            with self.assertRaisesRegex(ValueError,'Не удалось прочитать XLSX'):read_point_tables(path)


class TestV061ImportUI(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        previous.TestV06GeometryUI.setUpClass();cls.app=previous.TestV06GeometryUI.app
    setUp=previous.TestV06GeometryUI.setUp
    tearDown=previous.TestV06GeometryUI.tearDown
    record=previous.TestV06GeometryUI.record
    row=previous.TestV06GeometryUI.row
    handle=previous.TestV06GeometryUI.handle
    edit_number=previous.TestV06GeometryUI.edit_number

    def test_split_preserves_binding_on_second_half_and_undo(self):
        self.e.raw['objects'].append(dict(id='H1',type='hole',wall='outer',primitive_id='outer-1',z_mm=23.,radius_mm=300.,theta_deg=0.,parameters={},units={},provenance={'user_note':'keep'}))
        self.assertTrue(self.e.commit('Bound object'));before=project_to_json(self.w.cfg,self.w.nozzles);self.w.undo_stack.clear()
        self.e.primitive_table.setCurrentCell(self.row(self.e.primitive_table,'outer-1'),0);self.e.split_primitive()
        obj=next(o for o in self.w.cfg.geometry['objects'] if o['id']=='H1')
        self.assertTrue(obj['primitive_id'].startswith('НР'));self.assertNotEqual(obj['primitive_id'],'outer-1')
        self.assertEqual(obj['provenance'],{'user_note':'keep'});self.assertEqual(self.w.undo_stack.count(),1)
        self.w.undo_stack.undo();self.assertEqual(project_to_json(self.w.cfg,self.w.nozzles),before)

    def test_repair_invalid_radius_to_original_clears_draft(self):
        row=self.row(self.e.primitive_table,'reference-arc');original=self.e.primitive_table.item(row,9).text()
        self.e.primitive_table.item(row,9).setText('1');self.assertTrue(self.e._draft_invalid)
        self.e.primitive_table.item(row,9).setText(original);self.app.processEvents()
        self.assertFalse(self.e._draft_invalid);self.assertNotIn('Черновик',self.e.status.text())
        self.assertEqual(self.e.raw,self.w.cfg.geometry)

    def test_invalid_first_data_row_is_not_silently_treated_as_header(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'no_headers.xlsx'
            for first in ([0,-1],['NaN',250],[None,250],['=1+2',250]):
                book=Workbook();sheet=book.active;sheet.append(first);sheet.append([1,250]);book.save(path)
                dialog=PointImportDialog(path,self.e)
                self.assertEqual(dialog.first_row.value(),1);self.assertFalse(dialog.buttons.button(QDialogButtonBox.Ok).isEnabled())
                self.assertIn('Строка Excel 1',dialog.status.text());self.assertEqual(dialog.points,[]);dialog.close()

    def test_full_arc_parameter_paste_including_units_is_one_edit(self):
        row=self.row(self.e.primitive_table,'reference-arc');self.e.primitive_table.setCurrentCell(row,3)
        values=[.6,26.1,1.6,26.1,1.1,26.6,math.sqrt(.5),'','cm']
        self.e.primitive_table.paste_text('\t'.join(map(str,values)));self.app.processEvents()
        arc=Primitive(self.record('reference-arc'))
        self.assertEqual(arc.unit,'cm');self.assertAlmostEqual(arc.start[0],6.)
        self.assertAlmostEqual(arc.end[1],261.);self.assertAlmostEqual(arc.center[0],11.)
        self.assertIn(arc.id,self.e.paths);self.assertEqual(self.w.undo_stack.count(),1)
        self.assertEqual(self.w.cfg.geometry,self.e.raw)

    def test_import_preview_sheet_columns_unit_and_one_undo(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'точки.xlsx';workbook(path);dialog=PointImportDialog(path,self.e)
            dialog.sheet.setCurrentIndex(1);dialog.unit.setCurrentText('cm');dialog.wall.setCurrentIndex(1)
            self.assertEqual(dialog.points,[(2,5.,250.),(3,15.,260.)]);self.assertTrue(dialog.buttons.button(QDialogButtonBox.Ok).isEnabled())
            before=project_to_json(self.w.cfg,self.w.nozzles);old_ids={p['id'] for p in self.e.raw['primitives']}
            with patch('fuelfield.chamber_gui.QFileDialog.getOpenFileName',return_value=(str(path),'')),patch('fuelfield.chamber_gui.PointImportDialog',return_value=dialog),patch.object(dialog,'exec_',return_value=QDialog.Accepted):self.e.import_points()
            new=[p for p in self.e.raw['primitives'] if p['id'] not in old_ids]
            self.assertEqual([p['id'] for p in new],['ВН1','ВН2']);self.assertEqual(new[0]['start'],[5.,250.]);self.assertEqual(new[0]['provenance']['excel_row'],2)
            self.assertEqual(self.w.undo_stack.count(),1);self.w.undo_stack.undo();self.assertEqual(project_to_json(self.w.cfg,self.w.nozzles),before);self.w.undo_stack.redo();self.assertEqual(self.e.raw['primitives'][-1]['id'],'ВН2')
            dialog.close()

    def test_cancel_and_invalid_import_leave_project_unchanged(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'точки.xlsx';workbook(path);dialog=PointImportDialog(path,self.e);before=project_to_json(self.w.cfg,self.w.nozzles)
            dialog.r_column.setCurrentIndex(dialog.z_column.currentIndex());self.assertFalse(dialog.buttons.button(QDialogButtonBox.Ok).isEnabled());self.assertFalse(dialog.points)
            with patch('fuelfield.chamber_gui.QFileDialog.getOpenFileName',return_value=(str(path),'')),patch('fuelfield.chamber_gui.PointImportDialog',return_value=dialog),patch.object(dialog,'exec_',return_value=QDialog.Rejected):self.e.import_points()
            self.assertEqual(project_to_json(self.w.cfg,self.w.nozzles),before);self.assertEqual(self.w.undo_stack.count(),0);dialog.close()

    def test_moving_farthest_section_updates_outlet_preserving_retained_profile(self):
        before=deepcopy(self.e.raw['primitives']);objects=deepcopy(self.e.raw['objects']);outlet=self.e.raw['outlet_z_mm']
        row=next(i for i,s in enumerate(self.e.raw['sections']) if s['is_outlet']);identity=self.e.raw['sections'][row]['id']
        self.assertFalse(self.e.section_table.item(row,2).flags() & Qt.ItemIsUserCheckable)
        self.e.pages.setCurrentIndex(3);self.edit_number(self.e.section_table,identity,1,'21.375')
        self.assertEqual(self.e.raw['objects'],objects);self.assertEqual(self.e.raw['outlet_z_mm'],21.375)
        self.assertEqual(self.e.raw['primitives'][2:],before[2:])
        self.assertEqual(self.w.cfg.z_mm,[0.,21.375]);self.assertEqual(sum(s['is_outlet'] for s in self.e.raw['sections']),1)
        self.w.undo_stack.undo();self.assertEqual(self.w.cfg.z_mm,[0.,30.]);self.assertEqual(self.e.raw['primitives'],before)

    def test_sections_apply_even_when_variable_geometry_is_disabled(self):
        self.e.enabled.setChecked(False);self.app.processEvents();before=deepcopy(self.e.raw['primitives']);self.e.add_section()
        self.assertEqual(self.w.cfg.z_mm,[0.,30.,15.]);self.assertFalse(self.w.cfg.geometry['enabled']);self.assertEqual(self.e.raw['primitives'],before)

    def test_new_ids_are_readable_and_existing_ids_are_preserved(self):
        ids=[p['id'] for p in self.e.raw['primitives']];self.e.add_primitive('point');self.assertEqual(self.e.raw['primitives'][-1]['id'],'ОП1')
        self.e.add_primitive('point');self.assertEqual(self.e.raw['primitives'][-1]['id'],'ОП2')
        self.assertEqual([p['id'] for p in self.e.raw['primitives'][:len(ids)]],ids)
        self.e.add_object('hole');self.assertEqual(self.e.raw['objects'][-1]['id'],'H1')
        self.e.add_object('insert');self.assertEqual(self.e.raw['objects'][-1]['id'],'I1')
        self.e.add_section();self.assertEqual(self.e.raw['sections'][-1]['id'],'S3')

if __name__=='__main__':unittest.main()
